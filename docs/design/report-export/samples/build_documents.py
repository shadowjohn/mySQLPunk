"""Build small fictional schema specimens. Uses the Codex bundled Python runtime.

DOCX thumbnail PDFs are independent layout illustrations, not Office renders.
The PDF template thumbnails are rendered directly from the final PDF files.
"""
import json
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

ROOT = Path(__file__).resolve().parent
QA = ROOT / '_qa'
QA.mkdir(exist_ok=True)
DATA = json.loads((ROOT / 'demo-schema.json').read_text(encoding='utf-8'))
TABLES = DATA['tables']
FONT = 'Microsoft JhengHei'
pdfmetrics.registerFont(TTFont('JH', 'C:/Windows/Fonts/msjh.ttc', subfontIndex=0))
pdfmetrics.registerFont(TTFont('JHB', 'C:/Windows/Fonts/msjhbd.ttc', subfontIndex=0))
HEAD = ['欄位名稱', '欄位說明', '型態', '鍵', 'NULL', '預設值']

def shade(cell, color):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color); cell._tc.get_or_add_tcPr().append(el)

def set_font(run, size=9, bold=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if color: run.font.color.rgb = RGBColor.from_string(color)

def doc_table(doc, t, accent, classic_header=False):
    p = doc.add_paragraph(style='Heading 2'); p.add_run(t['name']+'  '+t['comment'])
    grid = doc.add_table(rows=1, cols=6); grid.autofit=False
    look = grid._tbl.tblPr.find(qn('w:tblLook'))
    if look is not None:
        # Keep the equivalent bitmask, compatible with the Office 2007 schema.
        for key in list(look.attrib):
            if key != qn('w:val'): del look.attrib[key]
    widths = [3.0,2.8,3.4,0.8,1.4,3.8]  # NULL needs room for four letters plus cell padding.
    for c,w in zip(grid.columns,widths): c.width=Cm(w)
    for cell,label,w in zip(grid.rows[0].cells,HEAD,widths):
        cell.width=Cm(w); shade(cell, accent); set_font(cell.paragraphs[0].add_run(label),9,True,'FFFFFF')
        if classic_header:
            margins=OxmlElement('w:tcMar'); left=OxmlElement('w:left')
            left.set(qn('w:w'),'75'); left.set(qn('w:type'),'dxa')  # 5 px at 96 dpi.
            margins.append(left); cell._tc.get_or_add_tcPr().append(margins)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat=OxmlElement('w:tblHeader'); grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    for n,row in enumerate(t['fields']):
        for cell,text,w in zip(grid.add_row().cells,row[:6],widths):
            cell.width=Cm(w); shade(cell, 'F2F5F9' if n%2==0 else 'FFFFFF')
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4)
            set_font(p.add_run(text),8.5)
        grid.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    details=['%s: %s' % (r[0],r[6]) for r in t['fields'] if r[6]!='-']
    p=doc.add_paragraph('其他屬性  '+'；'.join(details)); p.paragraph_format.space_after=Pt(8)
    for r in p.runs: set_font(r,8,color='536176')

def make_docx(name,title,delivery=False):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.5); sec.left_margin=Cm(1.5); sec.right_margin=Cm(1.5)
    for s in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[s].font.name=FONT; doc.styles[s]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    doc.styles['Normal'].font.size=Pt(9)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.size=Pt(24)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 2'].font.color.rgb=RGBColor.from_string('26446F' if delivery else '8C571A')
    doc.add_paragraph('demo_shop '+title,style='Title')
    doc.add_paragraph('樣版示例／示範結構')
    doc.add_paragraph('此範例呈現資料庫的資料表與欄位定義，供選擇報表樣版。demo_shop 為完全虛構的示範資料庫，包含 3 張資料表、16 個欄位，不含資料列。')
    if delivery:
        doc.add_paragraph('文件摘要',style='Heading 1')
        doc.add_paragraph('文件版本  1.0\n範例日期  2026-10-06\n資料庫名稱  demo_shop\n報表範圍  資料表、欄位、索引標記與關聯')
        doc.add_paragraph('資料表目錄',style='Heading 1')
        for n,t in enumerate(TABLES,1): doc.add_paragraph(f'{n:02d}  {t["name"]}  {t["comment"]}  {len(t["fields"])} 個欄位')
        doc.add_paragraph('關聯摘要',style='Heading 1')
        doc.add_paragraph('orders.customer_id 參照 customers.customer_id。\norder_items.order_id 參照 orders.order_id。')
        doc.add_page_break()
        doc.add_paragraph('資料表與欄位定義',style='Heading 1')
    for t in TABLES: doc_table(doc,t,'365C86' if delivery else 'B97B20',classic_header=not delivery)
    doc.add_paragraph('鍵標記：PK 主鍵；FK 外鍵；UK 唯一鍵。NULL 欄表示是否允許空值。「未設定」表示未宣告預設值。')
    doc.core_properties.title='demo_shop '+title
    doc.core_properties.author='mySQLPunk'
    doc.core_properties.subject='樣版示例／示範結構'
    doc.save(ROOT/name)

def p(text,size=9,color='#26374F',bold=False):
    return Paragraph(escape(text).replace('\n','<br/>'),ParagraphStyle('t',fontName='JHB' if bold else 'JH',fontSize=size,leading=size*1.45,textColor=colors.HexColor(color),spaceAfter=6,wordWrap='CJK'))

def pdf_table(t, width, accent, compact=False, classic_header=False):
    widths=[width*x for x in [.19,.17,.20,.06,.08,.30]]
    data=[[p(x,8,'#FFFFFF',True) for x in HEAD]]
    for row in t['fields']:
        data.append([p(x,8 if compact else 9) for x in row[:6]])
    table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor(accent)),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F1F5F9'),colors.white]),
        ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#C7D3E1'))]))
    if classic_header:
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,0),'MIDDLE'),('LEFTPADDING',(0,0),(-1,0),3.75)]))
    return table

def footer(canvas,doc):
    canvas.setFont('JH',8); canvas.setFillColor(colors.HexColor('#64748B'))
    canvas.drawString(doc.leftMargin,22,'mySQLPunk  /  demo_shop  /  樣版示例')
    canvas.drawRightString(doc.pagesize[0]-doc.rightMargin,22,str(doc.page))

def make_pdf(path,title,mode):
    compact=mode=='compact'; pagesize=A4
    doc=SimpleDocTemplate(str(path),pagesize=pagesize,leftMargin=35,rightMargin=35,topMargin=34,bottomMargin=36,title='demo_shop '+title,author='mySQLPunk')
    width=pagesize[0]-70; story=[p('demo_shop '+title,22,'#152743',True),p('樣版示例／示範結構',10,'#4C6B94')]
    if mode.startswith('docx'): story.append(p('設計示意縮圖，非 Office 實際轉檔畫面。',8,'#A06025'))
    story.append(p('完全虛構的資料庫定義。3 張資料表、16 個欄位，未包含資料列。',9))
    if mode=='docx-delivery':
        story += [Spacer(1,30),p('文件摘要',16,bold=True),p('文件版本  1.0\n範例日期  2026-10-06\n報表範圍  資料表、欄位、索引標記與關聯',11),Spacer(1,20),p('資料表目錄',16,bold=True)]
        for i,t in enumerate(TABLES,1): story.append(p(f'{i:02d}  {t["name"]}   {t["comment"]}   {len(t["fields"])} 個欄位',12))
        story += [Spacer(1,20),p('關聯摘要',16,bold=True),p('orders.customer_id 參照 customers.customer_id。\norder_items.order_id 參照 orders.order_id。',10)]
    else:
        for i,t in enumerate(TABLES):
            if mode=='standard' and i: story.append(PageBreak())
            story.append(Spacer(1,8 if compact else 16))
            story.append(p(f'{i+1:02d}  {t["name"]}   {t["comment"]}',12,'#315981',True))
            story.append(pdf_table(t,width,'#B97B20' if mode=='docx-classic' else '#365C86',compact or mode=='docx-classic',classic_header=mode=='docx-classic'))
            detail='；'.join(f'{r[0]}: {r[6]}' for r in t['fields'] if r[6]!='-')
            story.append(Spacer(1,5)); story.append(p('其他屬性  '+detail,8,'#536176'))
            if mode=='standard':
                story += [Spacer(1,20),p('結構摘要',12,bold=True),p(f'本表有 {len(t["fields"])} 個欄位。PK 為主鍵；FK 為外鍵；UK 為唯一鍵。NULL 欄顯示是否允許空值。',10)]
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

if __name__=='__main__':
    make_docx('docx-classic.docx','經典資料字典')
    make_docx('docx-delivery.docx','專案交付文件',True)
    make_pdf(ROOT/'pdf-standard.pdf','標準結構手冊','standard')
    make_pdf(ROOT/'pdf-compact.pdf','精簡審閱版','compact')
    make_pdf(QA/'docx-classic-illustration.pdf','經典資料字典','docx-classic')
    make_pdf(QA/'docx-delivery-illustration.pdf','專案交付文件','docx-delivery')
    print('Created 2 DOCX and 2 PDF specimens plus 2 labeled DOCX design illustrations.')
