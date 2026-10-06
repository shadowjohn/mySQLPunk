"""Validate specimens and publish small gallery thumbnails after visual inspection."""
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image

ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT/'demo-schema.json').read_text(encoding='utf-8'))
templates = [
    ('docx-classic','經典資料字典','docx','_qa/docx-classic-1.png','illustration',None),
    ('docx-delivery','專案交付文件','docx','_qa/docx-delivery-1.png','illustration',None),
    ('xlsx-tables','一表一工作表','xlsx','_qa/xlsx-tables-customers.png','artifact-tool-render',3),
    ('xlsx-inventory','欄位總覽清冊','xlsx','_qa/xlsx-inventory-field_inventory.png','artifact-tool-render',1),
    ('pdf-standard','標準結構手冊','pdf','_qa/pdf-standard-1.png','pdf-render',3),
    ('pdf-compact','精簡審閱版','pdf','_qa/pdf-compact-1.png','pdf-render',1)
]
out=[]
for stem,title,ext,preview,kind,count in templates:
    file=ROOT/f'{stem}.{ext}'
    if ext in ('docx','xlsx'):
        with ZipFile(file) as z: assert z.testzip() is None
    if ext=='docx':
        doc=Document(file)
        assert len(doc.tables)==3
        assert [len(t.rows)-1 for t in doc.tables]==[5,5,6]
        for table,expected in zip(doc.tables,DATA['tables']):
            assert [[c.text for c in r.cells] for r in table.rows[1:]]==[r[:6] for r in expected['fields']]
        text='\n'.join(p.text for p in doc.paragraphs)
        assert '樣版示例／示範結構' in text
        verified='OOXML ZIP 與 python-docx 重新讀取通過；3 表與 16 欄定義逐格吻合。Word 排版尚未驗收。'
    elif ext=='xlsx':
        wb=load_workbook(file,data_only=False)
        assert len(wb.worksheets)==count
        if stem=='xlsx-tables':
            for sheet,t in zip(wb.worksheets,DATA['tables']):
                assert sheet.title==t['name']
                assert [list(r) for r in sheet.iter_rows(min_row=7,max_row=6+len(t['fields']),max_col=7,values_only=True)]==t['fields']
                assert sheet.freeze_panes=='A7'
        else:
            expected=[[t['name'],t['comment'],*r] for t in DATA['tables'] for r in t['fields']]
            assert [list(r) for r in wb.active.iter_rows(min_row=7,max_row=22,max_col=9,values_only=True)]==expected
        verified='OOXML ZIP 與 openpyxl 重新讀取通過；16 欄定義逐格吻合；Artifact Tool 各工作表預覽已檢視。未啟動 Excel。'
    else:
        pdf=PdfReader(file); assert len(pdf.pages)==count
        text='\n'.join(p.extract_text() for p in pdf.pages)
        for t in DATA['tables']:
            assert t['name'] in text
            for r in t['fields']: assert r[0] in text
        verified=f'pypdf 重新讀取 {count} 頁；3 表與 16 欄名稱俱全；Poppler 最終 PDF 全頁影像已檢視。'
    img=Image.open(ROOT/preview).convert('RGB')
    img.thumbnail((1000,1000),Image.Resampling.LANCZOS)
    img.save(ROOT/f'{stem}.png',optimize=True)
    out.append({'id':stem,'name':title,'format':ext,'sample':f'{stem}.{ext}','thumbnail':f'{stem}.png','thumbnailKind':kind,'thumbnailNote':'設計示意，非 Office 實際轉檔畫面' if kind=='illustration' else ('Artifact Tool 工作表預覽，非 Excel 截圖' if ext=='xlsx' else '由最終 PDF 實際渲染'),'pageOrSheetCount':count,'bytes':file.stat().st_size,'verification':verified})
(ROOT/'manifest.json').write_text(json.dumps({'schema':'demo_shop','tables':3,'fields':16,'notice':DATA['notice'],'templates':out},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'validatedFiles':len(out),'tables':3,'fields':16,'templates':[x['id'] for x in out]},ensure_ascii=False))
