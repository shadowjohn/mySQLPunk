"""Create the fictional schema-specification DOCX specimen.

Use the Codex bundled Python runtime (python-docx). This specimen follows the
editable pale-blue pattern in the user reference, with the approved smaller
metadata-only scope. It is not output from the application report writer.
The original reference is never opened for writing or bundled with this file.

The thumbnail is produced separately from a read-only native Word PDF export.
The output path must not already exist, so running this builder cannot replace
another report. Example: python build-schema-specification-sample.py new.docx
"""

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
SCHEMA_PATH = ROOT / "docs/design/report-export/samples/demo-schema.json"
PAGE_WIDTH, PAGE_HEIGHT, MARGIN = 11906, 16838, 1134
FIELD_WIDTHS = [480, 2600, 2400, 450, 450, 2200, 1058]
INVENTORY_WIDTHS = [480, 4358, 4800]
SUMMARY_WIDTHS = [1800, 7838]
PALE_BLUE = "D9E1F2"
DEMO_INDEXES = {
    "customers": [("PRIMARY", ["customer_id"]), ("ux_customers_email", ["email"])],
    "orders": [("PRIMARY", ["order_id"]), ("ix_orders_customer_id", ["customer_id"])],
    "order_items": [("PRIMARY", ["item_id"]), ("ix_order_items_order_id", ["order_id"])],
}
# These indexes are explicit fictional fixture metadata, not inferred FK data.
# ux_customers_email is a UNIQUE index. The two ix_* indexes are non-unique.


def node(name, **attributes):
    element = OxmlElement("w:" + name)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    return element


def font(run, size=9, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    fonts.set(qn("w:cs"), "Calibri")


def paragraph_style(paragraph, *, before=0, after=60, center=False, keep=False):
    formatting = paragraph.paragraph_format
    formatting.space_before = Twips(before)
    formatting.space_after = Twips(after)
    formatting.line_spacing = 1
    formatting.keep_with_next = keep
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT


def paragraph(document, text, *, style=None, size=9, bold=False, before=0, after=60,
              new_page=False, keep=False):
    result = document.add_paragraph(style=style)
    paragraph_style(result, before=before, after=after, keep=keep)
    result.paragraph_format.page_break_before = new_page
    font(result.add_run(text), size, bold)
    return result


def table(document, widths):
    result = document.add_table(rows=0, cols=len(widths))
    result.autofit = False
    result.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = result._tbl.tblPr
    look = properties.find(qn("w:tblLook"))
    if look is not None:
        properties.remove(look)  # Explicit cell formatting; no Office 2010-only flags.
    properties.find(qn("w:tblW")).set(qn("w:w"), str(sum(widths)))
    properties.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    borders = node("tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(node(edge, val="single", sz=4, color="000000"))
    properties.insert_element_before(borders, "w:tblLayout", "w:tblCellMar", "w:tblLook")
    for column, width in zip(result.columns, widths):
        column.width = Twips(width)
    return result


def row(table_value, widths, values, *, header=False, label_column=False, centered=()):
    result = table_value.add_row()
    properties = result._tr.get_or_add_trPr()
    if header:
        properties.append(node("tblHeader"))
    for index, (cell, width, text) in enumerate(zip(result.cells, widths, values)):
        cell.width = Twips(width)
        cell_properties = cell._tc.get_or_add_tcPr()
        if header or (label_column and index == 0):
            cell_properties.append(node("shd", val="clear", fill=PALE_BLUE))
        margins = node("tcMar")
        for edge, amount in [("top", 60), ("left", 75), ("bottom", 60), ("right", 75)]:
            margins.append(node(edge, w=amount, type="dxa"))
        cell_properties.append(margins)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        paragraph_style(p, after=0, center=index in centered)
        font(p.add_run(str(text)), bold=header or (label_column and index == 0))
    return result


def split_type(value):
    """Match the native template's conservative character/decimal splitting."""
    match = re.fullmatch(
        r"(\s*(?:char|varchar|nchar|nvarchar|varchar2|nvarchar2|character(?:\s+varying)?|binary|varbinary)\s*)"
        r"\((\s*(?:[0-9]+(?:\s+(?:char|byte))?|max)\s*)\)(\s*)", value, re.I)
    if not match:
        match = re.fullmatch(
            r"(\s*(?:decimal|dec|numeric|number)\s*)"
            r"\((\s*[0-9]+(?:\s*,\s*-?[0-9]+)?\s*)\)((?:\s+(?:unsigned|zerofill))*\s*)", value, re.I)
    if not match:
        return value, "—"
    return match.group(1) + match.group(3), match.group(2)


def build(output):
    if output.exists():
        raise FileExistsError("Refusing to replace an existing report: " + str(output))
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Twips(PAGE_WIDTH), Twips(PAGE_HEIGHT)
    section.top_margin = section.bottom_margin = Twips(MARGIN)
    section.left_margin = section.right_margin = Twips(MARGIN)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.space_after = Twips(60)
    normal.paragraph_format.line_spacing = 1
    for name, size in [("Title", 18), ("Heading 1", 14)]:
        style = document.styles[name]
        style.font.name, style.font.size, style.font.bold = "Calibri", Pt(size), True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    document.core_properties.title = "系統資料庫規格書（示範）"
    document.core_properties.subject = "demo_shop fictional schema metadata specimen"
    document.core_properties.author = "mySQLPunk"
    document.core_properties.last_modified_by = "mySQLPunk"
    document.core_properties.comments = "Python-authored specimen, not application writer runtime output."

    paragraph(document, "系統資料庫規格書（示範）", style="Title", size=18, bold=True, after=180, keep=True)
    paragraph(document, "資料庫：" + schema["schema"])
    paragraph(document, "引擎：MySQL")
    paragraph(document, "作者：mySQLPunk 示範樣版")
    paragraph(document, "產生時間：2026-10-06 00:00")
    paragraph(document, "PK：主鍵；NN：不可為 NULL；V：是；-：否；?：未提供。長度／精度的 — 表示不適用或未提供。")
    paragraph(document, "資料表清單", style="Heading 1", size=14, bold=True,
              before=120, after=120, keep=True)
    inventory = table(document, INVENTORY_WIDTHS)
    row(inventory, INVENTORY_WIDTHS, ["No.", "資料表英文名稱", "資料表中文名稱"],
        header=True, centered=(0,))
    for index, item in enumerate(schema["tables"], 1):
        row(inventory, INVENTORY_WIDTHS, [index, item["name"], item["comment"]], centered=(0,))

    for index, item in enumerate(schema["tables"], 1):
        paragraph(document, str(index) + ". " + item["name"], style="Heading 1", size=14,
                  bold=True, before=120, after=120, new_page=True, keep=True)
        summary = table(document, SUMMARY_WIDTHS)
        pk = ", ".join(field[0] for field in item["fields"] if field[3] == "PK")
        indexes = "\n".join(name + " (" + ", ".join(columns) + ")"
                            for name, columns in DEMO_INDEXES[item["name"]])
        for label, value in [("資料表英文名稱", item["name"]),
                             ("資料表中文名稱", item["comment"]),
                             ("Primary Key", pk), ("Index", indexes)]:
            row(summary, SUMMARY_WIDTHS, [label, value], label_column=True)
        paragraph(document, "", after=60)
        fields = table(document, FIELD_WIDTHS)
        row(fields, FIELD_WIDTHS,
            ["No.", "欄位英文", "欄位中文", "PK", "NN", "資料型態", "長度／精度"],
            header=True, centered=(0, 3, 4, 6))
        for number, field in enumerate(item["fields"], 1):
            name, comment, original_type, key, nullable = field[:5]
            data_type, length = split_type(original_type)
            row(fields, FIELD_WIDTHS,
                [number, name, comment, "V" if key == "PK" else "-",
                 "V" if nullable == "NO" else "-", data_type, length],
                centered=(0, 3, 4, 6))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path, help="New DOCX path; existing files are rejected.")
    arguments = parser.parse_args()
    build(arguments.output.resolve())
